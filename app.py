import csv
import io
import re
import unicodedata
from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="Comparador Ergon x Intranet",
    page_icon="logo.png",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
<style>
:root { --navy:#16324F; --teal:#0F766E; --mint:#ECFDF5; --ink:#172033; --muted:#64748B; --amber:#B45309; }
.stApp { background: linear-gradient(180deg, #F8FAFC 0%, #FFFFFF 38%); color: var(--ink); }
[data-testid="stSidebar"] { background: linear-gradient(180deg, #102A43 0%, #16324F 100%); }
[data-testid="stSidebar"] * { color: #F8FAFC !important; }
[data-testid="stSidebar"] [data-testid="stFileUploader"] section { background: #FFFFFF !important; border: 1px solid #CBD5E1 !important; }
[data-testid="stSidebar"] [data-testid="stFileUploader"] section div { color: #16324F !important; }
[data-testid="stSidebar"] [data-testid="stFileUploader"] button { color: #16324F !important; background: #F8FAFC !important; border: 1px solid #94A3B8 !important; }
[data-testid="stSidebar"] [data-testid="stFileUploader"] button * { color: #16324F !important; }
[data-testid="stSidebar"] [data-testid="stFileUploader"] small { color: #CBD5E1 !important; }
.hero { padding: 26px 30px; border-radius: 20px; margin: 4px 0 24px; color: white;
    background: radial-gradient(circle at 90% 10%, #2DD4BF 0, transparent 30%), linear-gradient(120deg, #16324F 0%, #0F766E 100%);
    box-shadow: 0 12px 28px rgba(15,118,110,.18); }
.hero h1 { margin: 0; font-size: 2rem; letter-spacing: -.03em; }
.hero p { margin: 8px 0 0; color: #D9FDF5; font-size: .98rem; }
div[data-testid="stMetric"] { background: white; border: 1px solid #E2E8F0; border-radius: 14px; padding: 12px 14px; box-shadow: 0 5px 16px rgba(15,23,42,.05); }
div[data-testid="stMetricLabel"] { color: #64748B; }
div[data-testid="stMetricValue"] { color: #16324F; }
.section-kicker { text-transform: uppercase; letter-spacing: .12em; color: #0F766E; font-size: .72rem; font-weight: 700; margin-bottom: 4px; }
.stButton button, .stDownloadButton button { border-radius: 10px; font-weight: 650; }
</style>
""",
    unsafe_allow_html=True,
)

ACCEPTED_EXT = ["csv", "xls", "xlsx", "xlsm"]
ERGON_TO_INTRANET = {
    "Centro Int de Ass a Mul e a Cri D Regina S Campos": "CENTRO INTEGRADO DE ASSISTENCIA A MULHER E A CRIANCA DONA REGINA SIQUEIRA CAMPOS",
    "Hosp de Ref de Pedro Afonso - Leoncio de S Miranda": "HOSPITAL DE REFERENCIA DE PEDRO AFONSO - LEONCIO DE SOUSA MIRANDA",
    "Hosp Mat Infantil Edmunda a Cavalcante-tia Dede": "HOSPITAL MATERNO-INFANTIL EDMUNDA AIRES CAVALCANTE - TIA DEDE EM PORTO NACIONAL",
    "Hospital de Referencia de Alvorada do Tocantins": "HOSPITAL DE PEQUENO PORTE DE ALVORADA",
    "Hospital de Referencia de Araguaina": "HOSPITAL DE REFERENCIA DE ARAGUAINA",
    "Hospital de Referencia de Arraias": "HOSPITAL DE REFERENCIA DE ARRAIAS",
    "Hospital de Referencia de Augustinopolis": "HOSPITAL DE REFERENCIA DE AUGUSTINOPOLIS",
    "Hospital de Referencia de Dianopolis": "HOSPITAL DE REFERENCIA Dr. JAIMINHO",
    "Hospital de Referencia de Guarai": "HOSPITAL DE REFERENCIA DE GUARAI",
    "Hospital de Referencia de Gurupi": "HOSPITAL DE REFERENCIA DE GURUPI",
    "Hospital de Referencia de Miracema do Tocantins": "HOSPITAL DE REFERENCIA DE MIRACEMA DO TOCANTINS",
    "Hospital de Referencia de Porto Nacional": "HOSPITAL DE REFERENCIA DE PORTO NACIONAL",
    "Hospital de Referencia Tertuliano Corado Lustosa": "HOSPITAL DE REFERENCIA DE ARAGUACU",
    "Hospital e Maternidade Irmã Rita": "HOSPITAL E MATERNIDADE IRMA RITA",
    "Hospital Geral de Palmas Dr Francisco Ayres": "HOSPITAL GERAL DE PALMAS DR. FRANCISCO AYRES",
    "Hospital Regional Dr Alfredo Oliveira Barros": "HOSPITAL DE REFERENCIA DR. ALFREDO OLIVEIRA BARROS EM PARAISO DO TOCANTINS",
    "Hospital Regional Dr. João Lopes Machado": "HOSPITAL DE REFERENCIA DE XAMBIOA",
}

ERGON_ALIASES = {
    "numero funcional": "NUMFUNC",
    "numero do funcionario": "NUMFUNC",
    "numero vinculo": "NUMVINC",
    "nome servidor": "NOME",
    "lotacao": "LOTACAO",
    "lotacao atual": "LOTACAO",
    "situacao funcional": "SITUACAO",
}
INTRANET_ALIASES = {
    "numero funcional": "NUMFUNC",
    "numero do funcionario": "NUMFUNC",
    "numero vinculo": "NUMVINC",
    "nome servidor": "SERVIDOR",
    "ocupacao": "OCUPACAO",
    "descricao escala": "DESC. ESCALA",
    "desc escala": "DESC. ESCALA",
    "setor": "SETOR",
}


def safe_str(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return "" if text.casefold() in {"nan", "none", "nat"} else text


def header_key(value):
    text = safe_str(value).casefold()
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def canonicalize_columns(columns, aliases=None):
    aliases = aliases or {}
    used = set()
    result = []
    for raw in columns:
        label = re.sub(r"\s+", " ", safe_str(raw)).strip()
        canonical = aliases.get(header_key(label), label.upper())
        if canonical in used:
            suffix = 2
            candidate = f"{canonical}_{suffix}"
            while candidate in used:
                suffix += 1
            canonical = candidate
        used.add(canonical)
        result.append(canonical)
    return result


def detect_encoding(data):
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            data.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin1"


def read_text_robust(data):
    text = data.decode(detect_encoding(data), errors="replace").lstrip("\ufeff")
    if not text.strip():
        return pd.DataFrame()
    try:
        delimiter = csv.Sniffer().sniff(text[:10000], delimiters="\t;,|").delimiter
    except csv.Error:
        first = next((line for line in text.splitlines() if line.strip()), "")
        counts = {sep: first.count(sep) for sep in ("\t", ";", "|", ",")}
        delimiter = max(counts, key=counts.get) if max(counts.values()) else ","
    return pd.read_csv(io.StringIO(text), sep=delimiter, engine="python", header=None, dtype=object, keep_default_na=False, on_bad_lines="warn")


def read_excel_robust(data, kind):
    buffer = BytesIO(data)
    try:
        return pd.read_excel(buffer, engine="xlrd" if kind == "xls" else "openpyxl", header=None, dtype=object)
    except Exception:
        buffer.seek(0)
        return pd.read_excel(buffer, header=None, dtype=object)


@st.cache_data(show_spinner=False)
def read_any(data, filename):
    filename = filename.casefold()
    if filename.endswith(".csv"):
        return read_text_robust(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return read_excel_robust(data, "xls")
    if data[:4] == b"PK\x03\x04" or filename.endswith((".xlsx", ".xlsm")):
        return read_excel_robust(data, "xlsx")
    return read_text_robust(data)


def parse_raw(raw, aliases):
    raw = raw.dropna(how="all").reset_index(drop=True)
    if raw.empty:
        return pd.DataFrame()
    markers = {"NUMFUNC", "NUMVINC", "LOTACAO", "SETOR", "ESCALA", "CARGO"}
    header_index = 0
    for index in range(min(20, len(raw))):
        values = {header_key(value).upper().replace(" ", "") for value in raw.iloc[index].tolist()}
        if values & {marker.replace(" ", "") for marker in markers}:
            header_index = index
            break
    if raw.shape[1] <= 3:
        first = raw.iloc[:, 0].astype(str)
        header = first.iloc[header_index].split("\t")
        body = first.iloc[header_index + 1:].str.split("\t", expand=True)
        body = body.iloc[:, : len(header)]
        body.columns = canonicalize_columns(header, aliases)
        return body.reset_index(drop=True)
    body = raw.iloc[header_index + 1:].copy()
    body.columns = canonicalize_columns(raw.iloc[header_index].tolist(), aliases)
    return body.reset_index(drop=True)


def normalize_number(value):
    text = safe_str(value)
    if not text:
        return ""
    if re.fullmatch(r"-?\d+\.0+", text):
        text = text.split(".", 1)[0]
    digits = re.sub(r"\D", "", text)
    if not digits:
        return ""
    stripped = digits.lstrip("0")
    return stripped or "0"


def normalize_text(value):
    text = safe_str(value).casefold()
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text).strip().upper()


def make_key(numfunc, numvinc):
    functional = normalize_number(numfunc)
    vinculo = normalize_number(numvinc)
    return f"{functional}-{vinculo}" if functional and vinculo else ""


def clean_ergon(raw):
    df = raw.copy()
    for column in ("NUMFUNC", "NUMVINC", "NOME", "CPF", "CARGO", "LOTACAO", "SITUACAO"):
        if column not in df.columns:
            df[column] = ""
    for column in df.columns:
        df[column] = df[column].map(safe_str)
    df["NUMFUNC"] = df["NUMFUNC"].map(normalize_number)
    df["NUMVINC"] = df["NUMVINC"].map(normalize_number)
    df["CHAVE_VINCULO"] = [make_key(a, b) for a, b in zip(df["NUMFUNC"], df["NUMVINC"])]
    df["LOTACAO_NORM"] = df["LOTACAO"].map(normalize_text)
    hospital_map = {normalize_text(key): value for key, value in ERGON_TO_INTRANET.items()}
    df["HOSPITAL_ERGON"] = df["LOTACAO"]
    df["HOSPITAL_INTRANET"] = df["LOTACAO_NORM"].map(hospital_map).fillna("")
    df["HOSPITAL_INTRANET_NORM"] = df["HOSPITAL_INTRANET"].map(normalize_text)
    df = df[df["CHAVE_VINCULO"].ne("")].reset_index(drop=True)
    return df


def clean_intranet(raw):
    df = raw.copy()
    for column in ("NUMFUNC", "NUMVINC", "SERVIDOR", "CPF", "SETOR", "ESCALA", "CARGO", "OCUPACAO"):
        if column not in df.columns:
            df[column] = ""
    for column in df.columns:
        df[column] = df[column].map(safe_str)
    df["NUMFUNC"] = df["NUMFUNC"].map(normalize_number)
    df["NUMVINC"] = df["NUMVINC"].map(normalize_number)
    df["CHAVE_VINCULO"] = [make_key(a, b) for a, b in zip(df["NUMFUNC"], df["NUMVINC"])]
    df["HOSPITAL_INTRANET"] = df["SETOR"]
    df["HOSPITAL_INTRANET_NORM"] = df["SETOR"].map(normalize_text)
    df = df[df["CHAVE_VINCULO"].ne("")].reset_index(drop=True)
    return df


def validate_columns(df, required, label):
    missing = [column for column in required if column not in df.columns]
    if missing:
        return f"A planilha {label} não contém as colunas essenciais: {', '.join(missing)}."
    return None


def compare_datasets(ergon, intranet):
    mapped = ergon[ergon["HOSPITAL_INTRANET_NORM"].ne("")].copy()
    unmapped = ergon[ergon["HOSPITAL_INTRANET_NORM"].eq("")].copy()
    intranet_keys = set(zip(intranet["CHAVE_VINCULO"], intranet["HOSPITAL_INTRANET_NORM"]))
    mapped["ENCONTRADO_INTRANET"] = [
        (key, hospital) in intranet_keys
        for key, hospital in zip(mapped["CHAVE_VINCULO"], mapped["HOSPITAL_INTRANET_NORM"])
    ]
    mapped = mapped.drop_duplicates(subset=["CHAVE_VINCULO", "HOSPITAL_INTRANET_NORM"], keep="first")
    not_scaled = mapped[~mapped["ENCONTRADO_INTRANET"]].copy()
    scaled = mapped[mapped["ENCONTRADO_INTRANET"]].copy()
    return {
        "mapped": mapped.reset_index(drop=True),
        "not_scaled": not_scaled.reset_index(drop=True),
        "scaled": scaled.reset_index(drop=True),
        "unmapped": unmapped.drop_duplicates(subset=["CHAVE_VINCULO", "LOTACAO"], keep="first").reset_index(drop=True),
    }


def filter_results(df, hospital="(Todos)"):
    if hospital == "(Todos)":
        return df.copy().reset_index(drop=True)
    return df[df["HOSPITAL_ERGON"] == hospital].copy().reset_index(drop=True)


REPORT_COLUMNS = [
    "ORDEM", "CHAVE_VINCULO", "NOME", "NUMFUNC", "NUMVINC", "CPF",
    "EXERCICIO", "TIPO_VINCULO", "CARGO", "MOTIVO_DESATIVACAO", "PERIODO",
    "HOSPITAL_ERGON", "HOSPITAL_INTRANET", "ENCONTRADO_INTRANET",
]


def select_report_columns(df):
    """Retorna a saída principal na ordem solicitada, preenchendo campos ausentes."""
    output = df.copy()
    for column in REPORT_COLUMNS:
        if column not in output.columns:
            output[column] = ""
    return output[REPORT_COLUMNS].reset_index(drop=True)


def build_hospital_summary(mapped):
    if mapped.empty:
        return pd.DataFrame(columns=["HOSPITAL_ERGON", "HOSPITAL_INTRANET", "FOLHA_UNICOS", "ESCALADOS", "NAO_ESCALADOS"])
    return (
        mapped.groupby(["HOSPITAL_ERGON", "HOSPITAL_INTRANET"], dropna=False)
        .agg(
            FOLHA_UNICOS=("CHAVE_VINCULO", "nunique"),
            ESCALADOS=("ENCONTRADO_INTRANET", "sum"),
            NAO_ESCALADOS=("ENCONTRADO_INTRANET", lambda values: int((~values).sum())),
        )
        .reset_index()
        .sort_values(["NAO_ESCALADOS", "HOSPITAL_ERGON"], ascending=[False, True])
        .reset_index(drop=True)
    )


def format_br_int(value):
    try:
        return f"{int(value):,}".replace(",", ".")
    except (TypeError, ValueError):
        return str(value)


def format_dataframe(data):
    output = data.copy()
    for column in output.columns:
        if pd.api.types.is_numeric_dtype(output[column]):
            output[column] = output[column].map(format_br_int)
    return output


def display_dataframe(data, **kwargs):
    st.dataframe(format_dataframe(data), **kwargs)


def style_sheet(ws, df, color="1F4E78"):
    fill = PatternFill("solid", fgColor=color)
    font = Font(bold=True, color="FFFFFF", size=11)
    border_side = Side(border_style="thin", color="B7B7B7")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    for index, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=index)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        lengths = [len(str(name))]
        if len(df):
            lengths.append(int(df[name].astype(str).map(len).max()))
        ws.column_dimensions[get_column_letter(index)].width = min(max(max(lengths) + 2, 12), 55)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(df.columns)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    if len(df):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.print_area = f"A1:{get_column_letter(len(df.columns))}{ws.max_row}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.row_dimensions[1].height = 30


def prevent_excel_formula(value):
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def df_to_excel_bytes(sheets):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, data in sheets.items():
            safe_name = name[:31]
            prepared = data.copy()
            for column in prepared.columns:
                if pd.api.types.is_object_dtype(prepared[column]):
                    prepared[column] = prepared[column].map(prevent_excel_formula)
            prepared.to_excel(writer, sheet_name=safe_name, index=False)
            style_sheet(writer.sheets[safe_name], prepared)
    return output.getvalue()


def main():
    st.markdown(
        """
<div class="hero">
  <div class="section-kicker" style="color:#99F6E4">Auditoria de escalas</div>
  <h1>Comparador Ergon x Intranet</h1>
  <p>Encontre servidores na folha de pagamento que não estão escalados no hospital correspondente.</p>
</div>
""",
        unsafe_allow_html=True,
    )
    with st.sidebar:
        st.header("Arquivos de entrada")
        ergon_upload = st.file_uploader("Planilha do Ergon", type=ACCEPTED_EXT, key="ergon")
        intranet_upload = st.file_uploader("Planilha da Intranet", type=ACCEPTED_EXT, key="intranet")
        st.caption("A comparação usa a chave composta NUMFUNC + NUMVINC e considera o hospital correspondente.")
    if ergon_upload is None or intranet_upload is None:
        st.info("Envie as duas planilhas para iniciar a comparação.")
        st.stop()
    try:
        ergon_raw = parse_raw(read_any(ergon_upload.getvalue(), ergon_upload.name), ERGON_ALIASES)
        intranet_raw = parse_raw(read_any(intranet_upload.getvalue(), intranet_upload.name), INTRANET_ALIASES)
        ergon_error = validate_columns(ergon_raw, ["NUMFUNC", "NUMVINC", "LOTACAO"], "do Ergon")
        intranet_error = validate_columns(intranet_raw, ["NUMFUNC", "NUMVINC", "SETOR"], "da Intranet")
        if ergon_error or intranet_error:
            if ergon_error:
                st.error(ergon_error)
            if intranet_error:
                st.error(intranet_error)
            st.info("Confira se a primeira linha contém os cabeçalhos NUMFUNC, NUMVINC, LOTACAO e SETOR.")
            st.stop()
        ergon = clean_ergon(ergon_raw)
        intranet = clean_intranet(intranet_raw)
        comparison = compare_datasets(ergon, intranet)
    except Exception as exc:
        st.error(f"Não foi possível comparar os arquivos: {exc}")
        st.stop()
    mapped = comparison["mapped"]
    not_scaled = comparison["not_scaled"]
    unmapped = comparison["unmapped"]
    summary = build_hospital_summary(mapped)
    hospitals = ["(Todos)"] + sorted(not_scaled["HOSPITAL_ERGON"].dropna().unique().tolist()) if len(not_scaled) else ["(Todos)"]
    with st.sidebar:
        hospital = st.selectbox("Filtrar hospital", hospitals)
    filtered = filter_results(not_scaled, hospital)
    filtered_report = select_report_columns(filtered)
    st.success(f"Arquivos processados: {format_br_int(len(ergon))} registros do Ergon e {format_br_int(len(intranet))} vínculos/escala(s) da Intranet.")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Chaves na folha", format_br_int(mapped["CHAVE_VINCULO"].nunique()))
    k2.metric("Escalados", format_br_int(len(comparison["scaled"])))
    k3.metric("Não escalados", format_br_int(len(not_scaled)))
    k4.metric("Após filtro", format_br_int(len(filtered)))
    if len(unmapped):
        st.warning(f"{format_br_int(len(unmapped))} registro(s) do Ergon ficaram fora porque a lotação não está no mapeamento de hospitais.")
    tab_results, tab_summary, tab_unmapped = st.tabs(["Não escalados", "Resumo por hospital", "Fora do mapeamento"])
    with tab_results:
        st.subheader("Servidores na folha sem escala correspondente")
        st.caption("A chave é NUMFUNC-NUMVINC; o mesmo vínculo pode ter várias escalas na Intranet e basta uma ocorrência para ser considerado escalado.")
        if filtered.empty:
            st.success("Nenhum servidor não escalado encontrado para o filtro selecionado.")
        else:
            display_dataframe(filtered_report, use_container_width=True, hide_index=True, height=520)
    with tab_summary:
        st.subheader("Conferência por hospital")
        display_dataframe(summary, use_container_width=True, hide_index=True)
    with tab_unmapped:
        st.subheader("Lotações do Ergon sem correspondência")
        if unmapped.empty:
            st.success("Todas as lotações do Ergon consideradas estão mapeadas.")
        else:
            display_dataframe(unmapped[[c for c in ["CHAVE_VINCULO", "NOME", "NUMFUNC", "NUMVINC", "LOTACAO", "CARGO"] if c in unmapped.columns]], use_container_width=True, hide_index=True)
    st.markdown("---")
    st.subheader("Exportar comparação")
    st.caption("O Excel contém somente as 14 colunas do relatório, respeitando o hospital selecionado.")
    export = {"Nao escalados": filtered_report}
    st.download_button(
        "Baixar Excel da comparação",
        data=df_to_excel_bytes(export),
        file_name=f"comparacao_ergon_intranet_{datetime.today().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    main()
